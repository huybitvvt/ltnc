"""Build the handoff Word document from the authoritative Markdown files.
Requires python-docx. This script builds documentation only, not the application.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Ke hoach trien khai chi tiet - TroHub - Gemini.docx"
CHAPTERS = [
    ("Cách giao việc cho Gemini", ROOT / "START_HERE.md"),
    ("Kế hoạch tổng thể", ROOT / "GEMINI_IMPLEMENTATION_PLAN.md"),
    ("Quy tắc nghiệp vụ", ROOT / "docs/DOMAIN_RULES.md"),
    ("Backend, API và database", ROOT / "docs/BACKEND_DESIGN.md"),
    ("Frontend web và ứng dụng di động", ROOT / "docs/FRONTEND_DESIGN.md"),
    ("Các task triển khai", ROOT / "docs/IMPLEMENTATION_TASKS.md"),
    ("Kiểm thử và nghiệm thu", ROOT / "docs/TEST_ACCEPTANCE.md"),
    ("Các quyết định mặc định", ROOT / "docs/DECISIONS.md"),
    ("Nguồn tham khảo", ROOT / "docs/SOURCES.md"),
    ("Context GEMINI.md", ROOT / "GEMINI.md"),
    ("Checkpoint khởi đầu", ROOT / "docs/STATUS.md"),
]

def inline(paragraph, text):
    # Inline formatting is retained; Markdown links carry their URL as text in Word.
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    for token in re.split(r"(\*\*.+?\*\*|\x60[^\x60]+\x60)", text):
        if not token:
            continue
        run = paragraph.add_run(token[2:-2] if token.startswith("**") and token.endswith("**")
                                else token[1:-1] if token.startswith(chr(96)) and token.endswith(chr(96))
                                else token)
        if token.startswith("**") and token.endswith("**"):
            run.bold = True
        if token.startswith(chr(96)) and token.endswith(chr(96)):
            run.font.name = "Consolas"
            run.font.size = Pt(9)

def set_repeat_table_header(row):
    props = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    props.append(repeat)

def render(doc, source):
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    code = False
    skipped_title = False
    while i < len(lines):
        line = lines[i]
        if line.startswith(chr(96) * 3):
            code = not code
            i += 1
            continue
        if code:
            p = doc.add_paragraph(style="Code Block")
            p.add_run(line)
            i += 1
            continue
        if not line.strip() or re.fullmatch(r"\s*---+\s*", line):
            i += 1
            continue
        h = re.match(r"^(#{1,6})\s+(.+)$", line)
        if h:
            if len(h[1]) == 1 and not skipped_title:
                skipped_title = True
                i += 1
                continue
            p = doc.add_heading(level=min(4, max(2, len(h[1]))))
            inline(p, h[2])
            i += 1
            continue
        if line.lstrip().startswith("|"):
            raw_rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [v.strip() for v in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", v or " ") for v in cells):
                    raw_rows.append(cells)
                i += 1
            if not raw_rows:
                continue
            cols = max(len(r) for r in raw_rows)
            table = doc.add_table(rows=1, cols=cols)
            table.style = "Light Shading Accent 1"
            for row_index, values in enumerate(raw_rows):
                row = table.rows[0] if row_index == 0 else table.add_row()
                for col_index, value in enumerate(values):
                    inline(row.cells[col_index].paragraphs[0], value)
                    for p in row.cells[col_index].paragraphs:
                        p.paragraph_format.space_after = Pt(3)
                        for r in p.runs:
                            r.font.size = Pt(9)
                            if row_index == 0:
                                r.bold = True
                # Keep a single row together; headers repeat on subsequent pages.
                pr = row._tr.get_or_add_trPr()
                pr.append(OxmlElement("w:cantSplit"))
            set_repeat_table_header(table.rows[0])
            doc.add_paragraph()
            continue
        bullet = re.match(r"^\s*-\s+(.*)$", line)
        numbered = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            inline(p, bullet[1].replace("[ ]", "☐").replace("[x]", "☑"))
        elif numbered:
            p = doc.add_paragraph(style="List Number")
            inline(p, numbered[1])
        else:
            p = doc.add_paragraph()
            inline(p, line.lstrip("> "))
        i += 1

def main():
    for _, path in CHAPTERS:
        if not path.is_file():
            raise FileNotFoundError(path)
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [("Title", 27), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("183B56")
    from docx.enum.style import WD_STYLE_TYPE
    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1
    code.paragraph_format.left_indent = Cm(0.15)
    header = section.header.paragraphs[0]
    header.text = "TROHUB  |  KẾ HOẠCH TRIỂN KHAI VỚI GEMINI"
    header.style = "Caption"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Trang ")
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)

    doc.add_heading("TROHUB", 0)
    doc.add_heading("Kế hoạch triển khai chi tiết", 1)
    doc.add_paragraph("Hệ thống quản lý nhà trọ với hóa đơn minh bạch và cảnh báo tiêu thụ bất thường")
    doc.add_paragraph("BACKEND SPRING BOOT  •  WEB REACT  •  MOBILE REACT NATIVE")
    p = doc.add_paragraph()
    p.add_run("Ba vai trò trên cả web và app: OWNER · STAFF · TENANT").bold = True
    doc.add_paragraph("Phiên bản 2.0 | 21/09/2026")
    doc.add_paragraph("Tài liệu giao việc từ thiết kế đến nghiệm thu. Chưa có mã nguồn ứng dụng hoặc kết quả kiểm thử phần mềm.")
    doc.add_paragraph("Các file Markdown là nguồn để Gemini triển khai và cập nhật. Bản Word này là bản tổng hợp để đọc, không thay thế checkpoint đang làm.")
    doc.add_heading("Mục lục nội dung", 2)
    for idx, (title, _) in enumerate(CHAPTERS, 1):
        doc.add_paragraph(f"{idx}. {title}")
    for idx, (title, source) in enumerate(CHAPTERS, 1):
        doc.add_page_break()
        doc.add_heading(f"{idx}. {title}", 1)
        doc.add_paragraph(f"Nguồn: {source.relative_to(ROOT).as_posix()}", style="Caption")
        render(doc, source)
    doc.core_properties.title = "TroHub — Kế hoạch triển khai chi tiết với Gemini"
    doc.core_properties.subject = "Backend, web, mobile, nghiệp vụ, task và kiểm thử"
    doc.core_properties.author = "Project planning"
    doc.save(OUTPUT)
    print(OUTPUT)
    print(f"Chapters={len(CHAPTERS)}; paragraphs={len(doc.paragraphs)}; tables={len(doc.tables)}")

if __name__ == "__main__":
    main()

